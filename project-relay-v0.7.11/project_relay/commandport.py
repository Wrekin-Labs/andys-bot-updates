from __future__ import annotations

import fnmatch
import json
import os
import stat
import subprocess
import urllib.parse
import webbrowser
from pathlib import Path
from typing import Any

from .approvals import ApprovalStore
from .desktop_controls import desktop_context, assert_context
from .root_policy import configured_roots

FILESYSTEM_DEVICE_ID = "commandport:filesystem"
POWERSHELL_DEVICE_ID = "commandport:powershell"
DESKTOP_DEVICE_ID = "commandport:desktop"
MAX_TEXT_BYTES = 1_000_000
MAX_CAPTURE_CHARS = 128_000


def _allowed_roots() -> list[Path]:
    roots = [Path.home()]
    extra = os.getenv("PROJECT_RELAY_ALLOWED_ROOTS", "")
    for value in extra.split(os.pathsep):
        value = value.strip()
        if value:
            roots.append(Path(value).expanduser())
    roots.extend(configured_roots())

    unique: list[Path] = []
    seen: set[str] = set()
    for root in roots:
        resolved = root.resolve(strict=False)
        key = os.path.normcase(str(resolved))
        if key not in seen:
            seen.add(key)
            unique.append(resolved)
    return unique


def _inside(path: Path, root: Path) -> bool:
    try:
        common = os.path.commonpath([str(path), str(root)])
    except ValueError:
        return False
    return os.path.normcase(common) == os.path.normcase(str(root))


def _contains_link_or_reparse(path: Path) -> bool:
    """Reject symlink/junction traversal before resolving an approved path."""
    candidates = [path, *path.parents]
    for candidate in candidates:
        try:
            info = os.lstat(candidate)
        except (FileNotFoundError, OSError):
            continue
        if stat.S_ISLNK(info.st_mode):
            return True
        attrs = int(getattr(info, "st_file_attributes", 0) or 0)
        reparse = int(getattr(stat, "FILE_ATTRIBUTE_REPARSE_POINT", 0) or 0)
        if reparse and attrs & reparse:
            return True
    return False


def _resolve(path: str, *, must_exist: bool = False) -> Path:
    raw = str(path or "").strip()
    if not raw:
        raise ValueError("path is required")
    candidate = Path(raw).expanduser()
    if not candidate.is_absolute():
        candidate = Path.home() / candidate
    if _contains_link_or_reparse(candidate):
        raise PermissionError("Symbolic links and reparse-point traversal are not allowed")
    resolved = candidate.resolve(strict=False)

    if not any(_inside(resolved, root) for root in _allowed_roots()):
        raise PermissionError("Path is outside the approved Project Relay roots")
    if must_exist and not resolved.exists():
        raise FileNotFoundError(str(resolved))
    return resolved


def _entry(path: Path, root: Path) -> dict[str, Any]:
    try:
        stat = path.stat()
        size = stat.st_size
        mtime = stat.st_mtime
    except OSError:
        size = None
        mtime = None
    return {
        "name": path.name,
        "path": str(path),
        "relative": str(path.relative_to(root)) if path != root else ".",
        "type": "dir" if path.is_dir() else "file",
        "size": size,
        "mtime": mtime,
        "symlink": path.is_symlink(),
    }


def commandport_list_directory(path: str, depth: int = 1, max_entries: int = 500) -> dict[str, Any]:
    root = _resolve(path, must_exist=True)
    if not root.is_dir():
        raise NotADirectoryError(str(root))

    depth = max(0, min(int(depth), 4))
    max_entries = max(1, min(int(max_entries), 2000))
    results: list[dict[str, Any]] = []
    queue: list[tuple[Path, int]] = [(root, 0)]

    while queue and len(results) < max_entries:
        current, level = queue.pop(0)
        try:
            children = sorted(current.iterdir(), key=lambda p: (not p.is_dir(), p.name.lower()))
        except OSError:
            continue
        for child in children:
            if len(results) >= max_entries:
                break
            try:
                resolved = child.resolve(strict=False)
            except OSError:
                continue
            if not any(_inside(resolved, allowed) for allowed in _allowed_roots()):
                continue
            results.append(_entry(child, root))
            if level < depth and child.is_dir() and not child.is_symlink():
                queue.append((child, level + 1))

    return {
        "path": str(root),
        "count": len(results),
        "truncated": bool(queue) or len(results) >= max_entries,
        "items": results,
    }


def commandport_read_text_file(path: str, offset: int = 0, length: int = 64_000) -> dict[str, Any]:
    target = _resolve(path, must_exist=True)
    if not target.is_file():
        raise IsADirectoryError(str(target))
    offset = max(0, int(offset))
    length = max(1, min(int(length), 256_000))
    data = target.read_bytes()
    chunk = data[offset: offset + length]
    return {
        "path": str(target),
        "offset": offset,
        "length": len(chunk),
        "total_bytes": len(data),
        "text": chunk.decode("utf-8", errors="replace"),
        "has_more": offset + len(chunk) < len(data),
    }



def commandport_get_file_info(path: str) -> dict[str, Any]:
    target = _resolve(path, must_exist=True)
    info = target.stat()
    result: dict[str, Any] = {
        "path": str(target),
        "name": target.name,
        "type": "dir" if target.is_dir() else "file",
        "size": info.st_size,
        "mtime": info.st_mtime,
        "ctime": info.st_ctime,
        "symlink": target.is_symlink(),
        "suffix": target.suffix if target.is_file() else "",
        "mode": oct(info.st_mode & 0o777),
        "readable": os.access(target, os.R_OK),
        "writable": os.access(target, os.W_OK),
    }
    if not target.is_file():
        return result

    suffix = target.suffix.casefold()
    text_suffixes = {
        "", ".txt", ".md", ".py", ".js", ".ts", ".tsx", ".jsx", ".json",
        ".yaml", ".yml", ".toml", ".ini", ".cfg", ".csv", ".tsv", ".xml",
        ".html", ".css", ".log", ".ps1", ".bat", ".cmd", ".sql",
    }
    if info.st_size <= 2_000_000 and suffix in text_suffixes:
        try:
            data = target.read_bytes()
            if b"\x00" not in data[:8192]:
                line_count = 0 if not data else data.count(b"\n") + (0 if data.endswith(b"\n") else 1)
                result["line_count"] = line_count
                result["last_line"] = line_count - 1 if line_count else -1
                result["append_position"] = line_count
        except OSError:
            pass

    if suffix in {".xlsx", ".xlsm"} and info.st_size <= 50_000_000:
        try:
            from openpyxl import load_workbook
            workbook = load_workbook(str(target), read_only=True, data_only=False)
            try:
                result["sheets"] = [
                    {
                        "name": sheet.title,
                        "row_count": int(sheet.max_row or 0),
                        "column_count": int(sheet.max_column or 0),
                    }
                    for sheet in workbook.worksheets[:100]
                ]
            finally:
                workbook.close()
        except Exception:
            result["sheets"] = []
    return result


def commandport_search_files(
    path: str,
    pattern: str = "*",
    max_results: int = 200,
    max_depth: int = 8,
    include_dirs: bool = False,
) -> dict[str, Any]:
    root = _resolve(path, must_exist=True)
    pattern = str(pattern or "*").strip() or "*"
    if len(pattern) > 260:
        raise ValueError("pattern is too long")
    max_results = max(1, min(int(max_results), 1000))
    max_depth = max(0, min(int(max_depth), 32))
    allowed = _allowed_roots()
    matches: list[dict[str, Any]] = []
    scanned = 0
    truncated = False

    if root.is_file():
        if fnmatch.fnmatch(root.name.casefold(), pattern.casefold()):
            matches.append(_entry(root, root.parent))
        return {
            "path": str(root),
            "pattern": pattern,
            "count": len(matches),
            "scanned": 1,
            "truncated": False,
            "matches": matches,
        }
    if not root.is_dir():
        raise NotADirectoryError(str(root))

    for current_text, dirs, files in os.walk(root, topdown=True, followlinks=False):
        current = Path(current_text)
        try:
            level = len(current.relative_to(root).parts)
        except ValueError:
            dirs[:] = []
            continue

        safe_dirs: list[str] = []
        for name in dirs:
            candidate = current / name
            try:
                resolved = candidate.resolve(strict=False)
            except OSError:
                continue
            if candidate.is_symlink() or not any(_inside(resolved, item) for item in allowed):
                continue
            if level < max_depth:
                safe_dirs.append(name)
            if include_dirs:
                scanned += 1
                if fnmatch.fnmatch(name.casefold(), pattern.casefold()):
                    matches.append(_entry(candidate, root))
                    if len(matches) >= max_results:
                        truncated = True
                        break
        dirs[:] = safe_dirs
        if truncated:
            break

        for name in files:
            scanned += 1
            if scanned > 50_000:
                truncated = True
                break
            if not fnmatch.fnmatch(name.casefold(), pattern.casefold()):
                continue
            candidate = current / name
            try:
                resolved = candidate.resolve(strict=False)
            except OSError:
                continue
            if candidate.is_symlink() or not any(_inside(resolved, item) for item in allowed):
                continue
            matches.append(_entry(candidate, root))
            if len(matches) >= max_results:
                truncated = True
                break
        if truncated:
            break

    return {
        "path": str(root),
        "pattern": pattern,
        "count": len(matches),
        "scanned": scanned,
        "truncated": truncated,
        "matches": matches,
    }


def commandport_search_text(
    path: str,
    query: str,
    pattern: str = "*",
    case_sensitive: bool = False,
    max_results: int = 100,
    max_depth: int = 8,
    max_file_bytes: int = 1_000_000,
) -> dict[str, Any]:
    root = _resolve(path, must_exist=True)
    query = str(query or "")
    if not query:
        raise ValueError("query is required")
    if len(query) > 1000:
        raise ValueError("query is too long")
    pattern = str(pattern or "*").strip() or "*"
    max_results = max(1, min(int(max_results), 500))
    max_depth = max(0, min(int(max_depth), 32))
    max_file_bytes = max(1, min(int(max_file_bytes), 2_000_000))
    allowed = _allowed_roots()
    matches: list[dict[str, Any]] = []
    scanned_files = 0
    skipped_files = 0
    truncated = False
    needle = query if case_sensitive else query.casefold()

    def search_file(target: Path) -> None:
        nonlocal scanned_files, skipped_files, truncated
        if truncated:
            return
        try:
            if target.is_symlink() or not target.is_file():
                return
            resolved = target.resolve(strict=False)
            if not any(_inside(resolved, item) for item in allowed):
                return
            stat = target.stat()
            if stat.st_size > max_file_bytes:
                skipped_files += 1
                return
            data = target.read_bytes()
        except OSError:
            skipped_files += 1
            return
        scanned_files += 1
        if b"\x00" in data[:8192]:
            skipped_files += 1
            return
        text = data.decode("utf-8", errors="replace")
        for line_number, line in enumerate(text.splitlines(), start=1):
            haystack = line if case_sensitive else line.casefold()
            start = haystack.find(needle)
            if start < 0:
                continue
            preview = line.strip()
            if len(preview) > 320:
                left = max(0, start - 100)
                preview = line[left:left + 320].strip()
            matches.append({
                "path": str(target),
                "line": line_number,
                "column": start + 1,
                "preview": preview,
            })
            if len(matches) >= max_results:
                truncated = True
                return

    if root.is_file():
        if fnmatch.fnmatch(root.name.casefold(), pattern.casefold()):
            search_file(root)
    elif root.is_dir():
        for current_text, dirs, files in os.walk(root, topdown=True, followlinks=False):
            current = Path(current_text)
            try:
                level = len(current.relative_to(root).parts)
            except ValueError:
                dirs[:] = []
                continue
            safe_dirs: list[str] = []
            if level < max_depth:
                for name in dirs:
                    candidate = current / name
                    try:
                        resolved = candidate.resolve(strict=False)
                    except OSError:
                        continue
                    if not candidate.is_symlink() and any(_inside(resolved, item) for item in allowed):
                        safe_dirs.append(name)
            dirs[:] = safe_dirs

            for name in files:
                if scanned_files >= 5000:
                    truncated = True
                    break
                if not fnmatch.fnmatch(name.casefold(), pattern.casefold()):
                    continue
                search_file(current / name)
                if truncated:
                    break
            if truncated:
                break
    else:
        raise FileNotFoundError(str(root))

    return {
        "path": str(root),
        "query": query,
        "pattern": pattern,
        "case_sensitive": bool(case_sensitive),
        "count": len(matches),
        "scanned_files": scanned_files,
        "skipped_files": skipped_files,
        "truncated": truncated,
        "matches": matches,
    }


def commandport_list_windows(limit: int = 200) -> dict[str, Any]:
    if os.name != "nt":
        raise RuntimeError("CommandPort window listing is available on Windows only")
    limit = max(1, min(int(limit), 500))
    script = (
        "Get-Process | Where-Object { $_.MainWindowHandle -ne 0 -and $_.MainWindowTitle } | "
        "Sort-Object ProcessName | Select-Object -First " + str(limit) +
        " Id,ProcessName,MainWindowHandle,MainWindowTitle | ConvertTo-Json -Compress"
    )
    proc = subprocess.run(
        ["powershell", "-NoProfile", "-NonInteractive", "-Command", script],
        capture_output=True,
        text=True,
        timeout=20,
        check=False,
    )
    if proc.returncode != 0:
        raise RuntimeError(proc.stderr.strip() or "PowerShell window listing failed")
    raw = proc.stdout.strip()
    if not raw:
        rows: list[dict[str, Any]] = []
    else:
        parsed = json.loads(raw)
        rows = parsed if isinstance(parsed, list) else [parsed]
    return {"windows": rows, "count": len(rows)}


def commandport_list_processes(limit: int = 250) -> dict[str, Any]:
    if os.name != "nt":
        raise RuntimeError("CommandPort process listing is available on Windows only")
    limit = max(1, min(int(limit), 1000))
    script = (
        "Get-Process | Sort-Object ProcessName | "
        "Select-Object -First " + str(limit) + " Id,ProcessName,Path,CPU,WorkingSet64 | "
        "ConvertTo-Json -Compress"
    )
    proc = subprocess.run(
        ["powershell", "-NoProfile", "-NonInteractive", "-Command", script],
        capture_output=True,
        text=True,
        timeout=20,
        check=False,
    )
    if proc.returncode != 0:
        raise RuntimeError(proc.stderr.strip() or "PowerShell process listing failed")
    raw = proc.stdout.strip()
    if not raw:
        rows: list[dict[str, Any]] = []
    else:
        parsed = json.loads(raw)
        rows = parsed if isinstance(parsed, list) else [parsed]
    return {"processes": rows, "count": len(rows)}



def commandport_open_url(url: str) -> dict[str, Any]:
    """Open a normal http(s) URL in the workstation's default browser."""
    if os.name != "nt":
        raise RuntimeError("CommandPort browser navigation is available on Windows only")
    value = str(url or "").strip()
    if not value:
        raise ValueError("url is required")
    if len(value) > 4096:
        raise ValueError("url is too long")
    if any(ord(ch) < 32 for ch in value):
        raise ValueError("URL control characters are forbidden")
    parsed = urllib.parse.urlparse(value)
    if parsed.username is not None or parsed.password is not None:
        raise ValueError("Credentials in URLs are forbidden")
    # URL values are returned to the caller and appear in task records. Reject
    # credential-shaped query keys and all fragments before opening the browser.
    sensitive_keys = ("token", "secret", "password", "passwd", "api_key", "apikey",
                      "access_key", "auth", "credential", "session", "otp", "2fa",
                      "verification", "code", "card", "cvv", "cvc")
    if parsed.fragment or any(
        any(mark in key.casefold() for mark in sensitive_keys)
        for key, _ in urllib.parse.parse_qsl(parsed.query, keep_blank_values=True)
    ):
        raise ValueError("URLs with sensitive query parameters or fragments cannot be opened remotely")
    if parsed.scheme.lower() not in {"http", "https"} or not parsed.netloc:
        raise ValueError("Only normal http(s) URLs are allowed")
    opened = bool(webbrowser.open(value, new=2, autoraise=True))
    return {"url": value, "opened": opened}


def commandport_focus_window(window_handle: int) -> dict[str, Any]:
    """Bring one existing top-level Windows window to the foreground."""
    if os.name != "nt":
        raise RuntimeError("CommandPort window focus is available on Windows only")
    import ctypes

    hwnd = int(window_handle)
    if hwnd <= 0:
        raise ValueError("window_handle must be a positive integer")
    user32 = ctypes.windll.user32
    if not user32.IsWindow(hwnd):
        raise ValueError("window_handle does not identify a current window")
    user32.ShowWindow(hwnd, 9)  # SW_RESTORE
    ok = bool(user32.SetForegroundWindow(hwnd))
    return {"window_handle": hwnd, "focused": ok}


def commandport_capture_screen(
    max_width: int = 1024,
    jpeg_quality: int = 50,
) -> dict[str, Any]:
    """Fail closed: unrestricted desktop pixels cannot satisfy the privacy contract."""
    raise PermissionError(
        "Screen capture is disabled until a privacy-safe capture adapter is installed. "
        "No desktop pixels were captured."
    )


def _click_args(x: int, y: int, button: str, clicks: int) -> dict[str, Any]:
    if os.name != "nt":
        raise RuntimeError("CommandPort clicking is available on Windows only")
    button = str(button or "left").lower().strip()
    if button not in {"left", "right"}:
        raise ValueError("button must be left or right")
    if type(clicks) is not int or clicks not in (1, 2):
        raise ValueError("clicks must be 1 or 2")
    if type(x) is not int or type(y) is not int:
        raise ValueError("coordinates must be integers")
    context = desktop_context()
    left, top, right, bottom = context["rect"]
    if not (left <= x < right and top <= y < bottom):
        raise ValueError("click must be inside the active window")
    return {"x": x, "y": y, "button": button, "clicks": clicks, "context": context}


def commandport_prepare_click(
    x: int,
    y: int,
    button: str = "left",
    clicks: int = 1,
) -> dict[str, Any]:
    args = _click_args(x, y, button, clicks)
    record = ApprovalStore().create(
        "commandport_click",
        DESKTOP_DEVICE_ID,
        args,
        f"{args['button'].title()} click {args['clicks']}x at ({args['x']}, {args['y']})",
    )
    return {"approval": record.to_dict()}


def _perform_click(args: dict[str, Any]) -> dict[str, Any]:
    import ctypes
    import time

    user32 = ctypes.windll.user32
    assert_context(args["context"])
    if not user32.SetCursorPos(args["x"], args["y"]):
        raise RuntimeError("Windows could not position the cursor")
    if args["button"] == "right":
        down, up = 0x0008, 0x0010
    else:
        down, up = 0x0002, 0x0004
    for _ in range(args["clicks"]):
        assert_context(args["context"])
        user32.mouse_event(down, 0, 0, 0, 0)
        user32.mouse_event(up, 0, 0, 0, 0)
        time.sleep(0.08)
    return {"x": args["x"], "y": args["y"], "clicked": True}


def commandport_click(
    approval_id: str,
    x: int,
    y: int,
    button: str = "left",
    clicks: int = 1,
) -> dict[str, Any]:
    args = _click_args(x, y, button, clicks)
    ApprovalStore().consume(
        approval_id,
        action="commandport_click",
        device_id=DESKTOP_DEVICE_ID,
        arguments=args,
    )
    return _perform_click(args)


def commandport_owner_click(
    x: int,
    y: int,
    button: str = "left",
    clicks: int = 1,
) -> dict[str, Any]:
    raise PermissionError("Local approval is required; owner bypass is disabled")


def _type_args(text: str, press_enter: bool) -> dict[str, Any]:
    if os.name != "nt":
        raise RuntimeError("CommandPort typing is available on Windows only")
    value = str(text or "")
    if not value:
        raise ValueError("text is required")
    if len(value) > 4000:
        raise ValueError("text is too long")
    if type(press_enter) is not bool:
        raise ValueError("press_enter must be boolean")
    if any(ord(ch) < 32 or ord(ch) == 127 for ch in value):
        raise ValueError("Control characters are not allowed; use press_enter explicitly")
    return {"text": value, "press_enter": press_enter, "context": desktop_context()}


def commandport_prepare_type_text(
    text: str,
    press_enter: bool = False,
) -> dict[str, Any]:
    args = _type_args(text, press_enter)
    record = ApprovalStore().create(
        "commandport_type_text",
        DESKTOP_DEVICE_ID,
        args,
        "Type " + str(len(args["text"])) + " character(s) into " + args["context"]["application"] + " (window " + str(args["context"]["hwnd"]) + ")" +
        (" and press Enter" if args["press_enter"] else ""),
    )
    return {"approval": record.to_dict()}


def _perform_type_text(args: dict[str, Any]) -> dict[str, Any]:
    import ctypes
    from ctypes import wintypes

    user32 = ctypes.windll.user32
    INPUT_KEYBOARD = 1
    KEYEVENTF_KEYUP = 0x0002
    KEYEVENTF_UNICODE = 0x0004
    VK_RETURN = 0x0D

    ULONG_PTR = wintypes.WPARAM

    class MOUSEINPUT(ctypes.Structure):
        _fields_ = [
            ("dx", wintypes.LONG),
            ("dy", wintypes.LONG),
            ("mouseData", wintypes.DWORD),
            ("dwFlags", wintypes.DWORD),
            ("time", wintypes.DWORD),
            ("dwExtraInfo", ULONG_PTR),
        ]

    class KEYBDINPUT(ctypes.Structure):
        _fields_ = [
            ("wVk", wintypes.WORD),
            ("wScan", wintypes.WORD),
            ("dwFlags", wintypes.DWORD),
            ("time", wintypes.DWORD),
            ("dwExtraInfo", ULONG_PTR),
        ]

    class HARDWAREINPUT(ctypes.Structure):
        _fields_ = [
            ("uMsg", wintypes.DWORD),
            ("wParamL", wintypes.WORD),
            ("wParamH", wintypes.WORD),
        ]

    class INPUT_UNION(ctypes.Union):
        _fields_ = [
            ("mi", MOUSEINPUT),
            ("ki", KEYBDINPUT),
            ("hi", HARDWAREINPUT),
        ]

    class INPUT(ctypes.Structure):
        _anonymous_ = ("u",)
        _fields_ = [("type", wintypes.DWORD), ("u", INPUT_UNION)]

    def send_key(vk: int = 0, scan: int = 0, flags: int = 0) -> None:
        event = INPUT(
            type=INPUT_KEYBOARD,
            ki=KEYBDINPUT(vk, scan, flags, 0, 0),
        )
        sent = user32.SendInput(1, ctypes.byref(event), ctypes.sizeof(INPUT))
        if sent != 1:
            raise RuntimeError(
                f"Windows SendInput failed (INPUT size={ctypes.sizeof(INPUT)}, last_error={ctypes.get_last_error()})"
            )

    for ch in args["text"]:
        assert_context(args["context"])
        codepoint = ord(ch)
        if codepoint <= 0xFFFF:
            send_key(scan=codepoint, flags=KEYEVENTF_UNICODE)
            send_key(scan=codepoint, flags=KEYEVENTF_UNICODE | KEYEVENTF_KEYUP)
        else:
            codepoint -= 0x10000
            high = 0xD800 + (codepoint >> 10)
            low = 0xDC00 + (codepoint & 0x3FF)
            for unit in (high, low):
                send_key(scan=unit, flags=KEYEVENTF_UNICODE)
                send_key(scan=unit, flags=KEYEVENTF_UNICODE | KEYEVENTF_KEYUP)

    if args["press_enter"]:
        assert_context(args["context"])
        send_key(vk=VK_RETURN)
        send_key(vk=VK_RETURN, flags=KEYEVENTF_KEYUP)

    return {
        "characters_typed": len(args["text"]),
        "pressed_enter": args["press_enter"],
    }


def commandport_type_text(
    approval_id: str,
    text: str,
    press_enter: bool = False,
) -> dict[str, Any]:
    args = _type_args(text, press_enter)
    ApprovalStore().consume(
        approval_id,
        action="commandport_type_text",
        device_id=DESKTOP_DEVICE_ID,
        arguments=args,
    )
    return _perform_type_text(args)


def commandport_owner_type_text(
    text: str,
    press_enter: bool = False,
) -> dict[str, Any]:
    raise PermissionError("Local approval is required; owner bypass is disabled")


def _write_args(path: str, content: str, mode: str) -> tuple[Path, dict[str, Any]]:
    if mode not in {"rewrite", "append"}:
        raise ValueError("mode must be rewrite or append")
    encoded = content.encode("utf-8")
    if len(encoded) > MAX_TEXT_BYTES:
        raise ValueError("content exceeds the 1 MB CommandPort limit")
    target = _resolve(path, must_exist=False)
    args = {"path": str(target), "content": content, "mode": mode}
    return target, args


def commandport_prepare_write_file(path: str, content: str, mode: str = "rewrite") -> dict[str, Any]:
    target, args = _write_args(path, content, mode)
    record = ApprovalStore().create(
        "commandport_write_file",
        FILESYSTEM_DEVICE_ID,
        args,
        f"{mode.title()} UTF-8 file: {target}",
    )
    return {"approval": record.to_dict()}


def commandport_write_file(
    approval_id: str,
    path: str,
    content: str,
    mode: str = "rewrite",
) -> dict[str, Any]:
    target, args = _write_args(path, content, mode)
    ApprovalStore().consume(
        approval_id,
        action="commandport_write_file",
        device_id=FILESYSTEM_DEVICE_ID,
        arguments=args,
    )
    target.parent.mkdir(parents=True, exist_ok=True)
    file_mode = "a" if mode == "append" else "w"
    with target.open(file_mode, encoding="utf-8", newline="") as handle:
        handle.write(content)
    return {
        "path": str(target),
        "mode": mode,
        "bytes_written": len(content.encode("utf-8")),
    }


def _command_args(command: str, cwd: str | None, timeout_seconds: int) -> tuple[Path, dict[str, Any]]:
    if os.name != "nt":
        raise RuntimeError("CommandPort PowerShell execution is available on Windows only")
    command = str(command or "").strip()
    if not command:
        raise ValueError("command is required")
    if len(command) > 20_000:
        raise ValueError("command is too long")
    timeout_seconds = max(1, min(int(timeout_seconds), 300))
    workdir = _resolve(cwd or str(Path.home()), must_exist=True)
    if not workdir.is_dir():
        raise NotADirectoryError(str(workdir))
    args = {
        "command": command,
        "cwd": str(workdir),
        "timeout_seconds": timeout_seconds,
    }
    return workdir, args


def commandport_prepare_run_command(
    command: str,
    cwd: str | None = None,
    timeout_seconds: int = 60,
) -> dict[str, Any]:
    workdir, args = _command_args(command, cwd, timeout_seconds)
    preview = command.replace("\r", " ").replace("\n", " ")[:180]
    record = ApprovalStore().create(
        "commandport_run_command",
        POWERSHELL_DEVICE_ID,
        args,
        f"Run PowerShell in {workdir}: {preview}",
    )
    return {"approval": record.to_dict()}


def commandport_run_command(
    approval_id: str,
    command: str,
    cwd: str | None = None,
    timeout_seconds: int = 60,
) -> dict[str, Any]:
    workdir, args = _command_args(command, cwd, timeout_seconds)
    ApprovalStore().consume(
        approval_id,
        action="commandport_run_command",
        device_id=POWERSHELL_DEVICE_ID,
        arguments=args,
    )
    try:
        proc = subprocess.run(
            ["powershell", "-NoProfile", "-NonInteractive", "-Command", command],
            cwd=str(workdir),
            capture_output=True,
            text=True,
            timeout=args["timeout_seconds"],
            check=False,
        )
        return {
            "cwd": str(workdir),
            "returncode": proc.returncode,
            "stdout": proc.stdout[-MAX_CAPTURE_CHARS:],
            "stderr": proc.stderr[-MAX_CAPTURE_CHARS:],
            "timed_out": False,
        }
    except subprocess.TimeoutExpired as exc:
        stdout = exc.stdout or ""
        stderr = exc.stderr or ""
        if isinstance(stdout, bytes):
            stdout = stdout.decode("utf-8", errors="replace")
        if isinstance(stderr, bytes):
            stderr = stderr.decode("utf-8", errors="replace")
        return {
            "cwd": str(workdir),
            "returncode": None,
            "stdout": str(stdout)[-MAX_CAPTURE_CHARS:],
            "stderr": str(stderr)[-MAX_CAPTURE_CHARS:],
            "timed_out": True,
        }


def commandport_owner_run_command(
    command: str,
    cwd: str | None = None,
    timeout_seconds: int = 60,
) -> dict[str, Any]:
    """Run the existing bounded PowerShell adapter when local Owner Full Control is enabled."""
    from .owner_full_control import require_enabled
    require_enabled()
    workdir, args = _command_args(command, cwd, timeout_seconds)
    try:
        proc = subprocess.run(
            ["powershell", "-NoProfile", "-NonInteractive", "-Command", command],
            cwd=str(workdir), capture_output=True, text=True,
            timeout=args["timeout_seconds"], check=False,
        )
        return {
            "cwd": str(workdir), "returncode": proc.returncode,
            "stdout": proc.stdout[-MAX_CAPTURE_CHARS:],
            "stderr": proc.stderr[-MAX_CAPTURE_CHARS:], "timed_out": False,
        }
    except subprocess.TimeoutExpired as exc:
        stdout = exc.stdout or ""; stderr = exc.stderr or ""
        if isinstance(stdout, bytes): stdout = stdout.decode("utf-8", errors="replace")
        if isinstance(stderr, bytes): stderr = stderr.decode("utf-8", errors="replace")
        return {
            "cwd": str(workdir), "returncode": None,
            "stdout": str(stdout)[-MAX_CAPTURE_CHARS:],
            "stderr": str(stderr)[-MAX_CAPTURE_CHARS:], "timed_out": True,
        }
