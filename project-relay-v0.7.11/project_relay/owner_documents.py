from __future__ import annotations

import csv
import hashlib
import io
import json
import os
import tempfile
import time
import uuid
from pathlib import Path
from typing import Any

from .commandport import _resolve
from .owner_full_control import require_enabled
from .state import state_dir

MAX_DOCUMENT_BYTES = 10_000_000
MAX_ROWS = 5_000
MAX_COLS = 200
MAX_CELLS = 2_000
MAX_TEXT_CHARS = 500_000


def _sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _target(path: str) -> Path:
    target = _resolve(path, must_exist=False)
    if target.exists() and target.is_dir():
        raise IsADirectoryError(str(target))
    return target


def _checkpoint(target: Path, before: bytes, after: bytes) -> str | None:
    if not before:
        return None
    root = state_dir() / "document-checkpoints"
    root.mkdir(parents=True, exist_ok=True)
    checkpoint_id = "doccp-" + uuid.uuid4().hex
    (root / f"{checkpoint_id}.bak").write_bytes(before)
    (root / f"{checkpoint_id}.json").write_text(
        json.dumps(
            {
                "checkpoint_id": checkpoint_id,
                "target": str(target),
                "before_sha256": _sha256(before),
                "after_sha256": _sha256(after),
                "created_at": time.time(),
            },
            separators=(",", ":"),
            sort_keys=True,
        ),
        encoding="utf-8",
    )
    return checkpoint_id


def _commit_bytes(
    target: Path,
    data: bytes,
    expected_sha256: str | None = None,
) -> dict[str, Any]:
    if len(data) > MAX_DOCUMENT_BYTES:
        raise ValueError("document output exceeds size limit")
    before = target.read_bytes() if target.is_file() else b""
    before_sha = _sha256(before) if before else None
    if expected_sha256 is not None:
        if not before:
            raise RuntimeError("expected_sha256 was provided but target does not exist")
        if before_sha != str(expected_sha256).lower():
            raise RuntimeError("document changed since preview/read; write aborted")
    target.parent.mkdir(parents=True, exist_ok=True)
    temp_name: str | None = None
    try:
        with tempfile.NamedTemporaryFile(
            mode="wb", delete=False, dir=str(target.parent), prefix=f".{target.name}.relay-"
        ) as handle:
            temp_name = handle.name
            handle.write(data)
            handle.flush()
            os.fsync(handle.fileno())
        os.replace(temp_name, target)
        temp_name = None
    finally:
        if temp_name:
            Path(temp_name).unlink(missing_ok=True)
    checkpoint_id = _checkpoint(target, before, data)
    return {
        "path": str(target),
        "bytes_written": len(data),
        "before_sha256": before_sha,
        "sha256": _sha256(data),
        **({"checkpoint_id": checkpoint_id} if checkpoint_id else {}),
    }


def owner_write_csv(
    path: str,
    rows: list[list[Any]],
    expected_sha256: str | None = None,
) -> dict[str, Any]:
    require_enabled()
    if not isinstance(rows, list) or len(rows) > MAX_ROWS:
        raise ValueError("rows must be a list within the row limit")
    out = io.StringIO(newline="")
    writer = csv.writer(out)
    for row in rows:
        if not isinstance(row, list) or len(row) > MAX_COLS:
            raise ValueError("each row must be a list within the column limit")
        writer.writerow(["" if value is None else value for value in row])
    data = out.getvalue().encode("utf-8")
    return {"kind": "csv", **_commit_bytes(_target(path), data, expected_sha256)}


def owner_write_xlsx(
    path: str,
    sheets: dict[str, list[list[Any]]],
    expected_sha256: str | None = None,
) -> dict[str, Any]:
    require_enabled()
    if not isinstance(sheets, dict) or not sheets or len(sheets) > 50:
        raise ValueError("sheets must contain 1 to 50 worksheets")
    from openpyxl import Workbook

    wb = Workbook()
    default = wb.active
    wb.remove(default)
    for name, rows in sheets.items():
        title = str(name)
        if not title or len(title) > 31:
            raise ValueError("worksheet names must contain 1 to 31 characters")
        if not isinstance(rows, list) or len(rows) > MAX_ROWS:
            raise ValueError("worksheet row limit exceeded")
        ws = wb.create_sheet(title)
        for row in rows:
            if not isinstance(row, list) or len(row) > MAX_COLS:
                raise ValueError("worksheet column limit exceeded")
            ws.append(row)
    out = io.BytesIO()
    wb.save(out)
    wb.close()
    return {"kind": "xlsx", "sheet_count": len(sheets), **_commit_bytes(_target(path), out.getvalue(), expected_sha256)}


def owner_update_xlsx_cells(
    path: str,
    sheet: str,
    cells: dict[str, Any],
    expected_sha256: str | None = None,
) -> dict[str, Any]:
    require_enabled()
    target = _resolve(path, must_exist=True)
    if target.suffix.casefold() not in {".xlsx", ".xlsm"}:
        raise ValueError("target must be XLSX/XLSM")
    if not isinstance(cells, dict) or not cells or len(cells) > MAX_CELLS:
        raise ValueError("cells must contain 1 to 2000 coordinate/value pairs")
    before = target.read_bytes()
    before_sha = _sha256(before)
    if expected_sha256 and before_sha != str(expected_sha256).lower():
        raise RuntimeError("document changed since preview/read; write aborted")
    from openpyxl import load_workbook

    keep_vba = target.suffix.casefold() == ".xlsm"
    wb = load_workbook(filename=io.BytesIO(before), keep_vba=keep_vba)
    try:
        if sheet not in wb.sheetnames:
            raise KeyError(f"worksheet not found: {sheet}")
        ws = wb[sheet]
        for coordinate, value in cells.items():
            key = str(coordinate).upper().strip()
            if not key or len(key) > 16:
                raise ValueError("invalid cell coordinate")
            ws[key] = value
        out = io.BytesIO()
        wb.save(out)
    finally:
        wb.close()
    result = _commit_bytes(target, out.getvalue(), before_sha)
    return {"kind": "xlsx", "sheet": sheet, "cells_updated": len(cells), **result}


def owner_update_xlsx_range(
    path: str,
    range_ref: str,
    values: list[list[Any]],
    expected_sha256: str | None = None,
) -> dict[str, Any]:
    require_enabled()
    target = _resolve(path, must_exist=True)
    if target.suffix.casefold() not in {".xlsx", ".xlsm"}:
        raise ValueError("target must be XLSX/XLSM")
    reference = str(range_ref or "").strip()
    if not reference or len(reference) > 200:
        raise ValueError("range_ref is required")
    if not isinstance(values, list) or not values or len(values) > MAX_ROWS:
        raise ValueError("values must contain 1 to 5000 rows")
    if any(not isinstance(row, list) or len(row) > MAX_COLS for row in values):
        raise ValueError("each values row must be a list within the column limit")

    before = target.read_bytes()
    before_sha = _sha256(before)
    if expected_sha256 and before_sha != str(expected_sha256).lower():
        raise RuntimeError("document changed since preview/read; write aborted")

    from openpyxl import load_workbook
    from openpyxl.utils.cell import range_to_tuple

    try:
        sheet_name, bounds = range_to_tuple(reference)
    except ValueError as exc:
        raise ValueError("range_ref must use Sheet!A1:C10 form") from exc
    min_col, min_row, max_col, max_row = bounds
    expected_rows = max_row - min_row + 1
    expected_cols = max_col - min_col + 1
    if len(values) != expected_rows or any(len(row) != expected_cols for row in values):
        raise ValueError(
            f"values dimensions must exactly match range: {expected_rows}x{expected_cols}"
        )
    if expected_rows * expected_cols > MAX_CELLS:
        raise ValueError("range exceeds 2000-cell update limit")

    keep_vba = target.suffix.casefold() == ".xlsm"
    wb = load_workbook(filename=io.BytesIO(before), keep_vba=keep_vba)
    try:
        if sheet_name not in wb.sheetnames:
            raise KeyError(f"worksheet not found: {sheet_name}")
        ws = wb[sheet_name]
        for row_offset, row_values in enumerate(values):
            for col_offset, value in enumerate(row_values):
                ws.cell(
                    row=min_row + row_offset,
                    column=min_col + col_offset,
                    value=value,
                )
        out = io.BytesIO()
        wb.save(out)
    finally:
        wb.close()

    result = _commit_bytes(target, out.getvalue(), before_sha)
    return {
        "kind": "xlsx",
        "range": reference,
        "rows_updated": expected_rows,
        "columns_updated": expected_cols,
        "cells_updated": expected_rows * expected_cols,
        **result,
    }


def owner_write_docx(
    path: str,
    text: str,
    title: str | None = None,
    expected_sha256: str | None = None,
) -> dict[str, Any]:
    require_enabled()
    value = str(text)
    if len(value) > MAX_TEXT_CHARS:
        raise ValueError("document text exceeds limit")
    from docx import Document

    doc = Document()
    if title:
        doc.add_heading(str(title)[:240], level=1)
    for paragraph in value.split("\n\n"):
        doc.add_paragraph(paragraph)
    out = io.BytesIO()
    doc.save(out)
    return {"kind": "docx", **_commit_bytes(_target(path), out.getvalue(), expected_sha256)}


def owner_write_pdf(
    path: str,
    text: str,
    title: str | None = None,
    expected_sha256: str | None = None,
) -> dict[str, Any]:
    require_enabled()
    value = str(text)
    if len(value) > MAX_TEXT_CHARS:
        raise ValueError("document text exceeds limit")
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    out = io.BytesIO()
    pdf = canvas.Canvas(out, pagesize=A4)
    width, height = A4
    y = height - 54
    if title:
        pdf.setFont("Helvetica-Bold", 14)
        pdf.drawString(54, y, str(title)[:100])
        y -= 28
    pdf.setFont("Helvetica", 10)
    for source_line in value.splitlines() or [""]:
        line = source_line
        chunks = [line[i:i + 100] for i in range(0, max(1, len(line)), 100)] or [""]
        for chunk in chunks:
            if y < 54:
                pdf.showPage()
                pdf.setFont("Helvetica", 10)
                y = height - 54
            pdf.drawString(54, y, chunk)
            y -= 14
    pdf.save()
    return {"kind": "pdf", **_commit_bytes(_target(path), out.getvalue(), expected_sha256)}


def owner_rollback_document(checkpoint_id: str) -> dict[str, Any]:
    require_enabled()
    value = str(checkpoint_id or "").strip()
    if not value.startswith("doccp-") or len(value) != 38:
        raise ValueError("invalid document checkpoint id")
    root = state_dir() / "document-checkpoints"
    meta_path = root / f"{value}.json"
    backup_path = root / f"{value}.bak"
    if not meta_path.is_file() or not backup_path.is_file():
        raise KeyError("document checkpoint not found")
    meta = json.loads(meta_path.read_text(encoding="utf-8"))
    target = _resolve(str(meta["target"]), must_exist=True)
    current = target.read_bytes()
    if _sha256(current) != str(meta["after_sha256"]):
        raise RuntimeError("document changed after checkpoint; rollback refused")
    before = backup_path.read_bytes()
    if _sha256(before) != str(meta["before_sha256"]):
        raise RuntimeError("document checkpoint integrity check failed")
    result = _commit_bytes(target, before, str(meta["after_sha256"]))
    return {"rolled_back": True, "checkpoint_id": value, **result}


def owner_write_json(
    path: str,
    value: Any,
    pretty: bool = True,
    expected_sha256: str | None = None,
) -> dict[str, Any]:
    require_enabled()
    text = json.dumps(value, indent=2 if pretty else None, ensure_ascii=False, sort_keys=bool(pretty))
    data = (text + ("\n" if pretty else "")).encode("utf-8")
    return {"kind": "json", **_commit_bytes(_target(path), data, expected_sha256)}


def owner_replace_docx_text(
    path: str,
    find: str,
    replace: str,
    expected_sha256: str | None = None,
) -> dict[str, Any]:
    require_enabled()
    if not find:
        raise ValueError("find text is required")
    target = _resolve(path, must_exist=True)
    if target.suffix.casefold() != ".docx":
        raise ValueError("target must be DOCX")
    before = target.read_bytes()
    before_sha = _sha256(before)
    if expected_sha256 and before_sha != str(expected_sha256).lower():
        raise RuntimeError("document changed since preview/read; write aborted")
    from docx import Document

    doc = Document(io.BytesIO(before))
    replacements = 0

    def apply_paragraph(paragraph) -> None:
        nonlocal replacements
        count = paragraph.text.count(find)
        if count:
            paragraph.text = paragraph.text.replace(find, replace)
            replacements += count

    for paragraph in doc.paragraphs:
        apply_paragraph(paragraph)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    apply_paragraph(paragraph)
    if replacements == 0:
        return {"kind": "docx", "changed": False, "replacements": 0, "path": str(target), "sha256": before_sha}
    out = io.BytesIO()
    doc.save(out)
    result = _commit_bytes(target, out.getvalue(), before_sha)
    return {"kind": "docx", "changed": True, "replacements": replacements, **result}


def owner_pdf_delete_pages(
    path: str,
    pages: list[int],
    expected_sha256: str | None = None,
) -> dict[str, Any]:
    require_enabled()
    target = _resolve(path, must_exist=True)
    if target.suffix.casefold() != ".pdf":
        raise ValueError("target must be PDF")
    if not isinstance(pages, list) or not pages or len(pages) > 500:
        raise ValueError("pages must contain 1 to 500 one-based page numbers")
    before = target.read_bytes()
    before_sha = _sha256(before)
    if expected_sha256 and before_sha != str(expected_sha256).lower():
        raise RuntimeError("document changed since preview/read; write aborted")
    from pypdf import PdfReader, PdfWriter

    reader = PdfReader(io.BytesIO(before))
    remove = {int(page) for page in pages}
    if any(page < 1 or page > len(reader.pages) for page in remove):
        raise ValueError("page number is outside the document")
    if len(remove) >= len(reader.pages):
        raise ValueError("refusing to delete every PDF page")
    writer = PdfWriter()
    for index, page in enumerate(reader.pages, start=1):
        if index not in remove:
            writer.add_page(page)
    out = io.BytesIO()
    writer.write(out)
    result = _commit_bytes(target, out.getvalue(), before_sha)
    return {
        "kind": "pdf",
        "deleted_pages": sorted(remove),
        "page_count": len(reader.pages) - len(remove),
        **result,
    }


def owner_pdf_insert_pdf(
    path: str,
    source_path: str,
    position: int = 0,
    expected_sha256: str | None = None,
) -> dict[str, Any]:
    require_enabled()
    target = _resolve(path, must_exist=True)
    source = _resolve(source_path, must_exist=True)
    if target.suffix.casefold() != ".pdf" or source.suffix.casefold() != ".pdf":
        raise ValueError("target and source must be PDF")
    before = target.read_bytes()
    before_sha = _sha256(before)
    if expected_sha256 and before_sha != str(expected_sha256).lower():
        raise RuntimeError("document changed since preview/read; write aborted")
    from pypdf import PdfReader, PdfWriter

    target_reader = PdfReader(io.BytesIO(before))
    source_reader = PdfReader(str(source))
    insert_at = max(0, min(int(position), len(target_reader.pages)))
    writer = PdfWriter()
    for page in target_reader.pages[:insert_at]:
        writer.add_page(page)
    for page in source_reader.pages:
        writer.add_page(page)
    for page in target_reader.pages[insert_at:]:
        writer.add_page(page)
    out = io.BytesIO()
    writer.write(out)
    result = _commit_bytes(target, out.getvalue(), before_sha)
    return {
        "kind": "pdf",
        "inserted_pages": len(source_reader.pages),
        "position": insert_at,
        "page_count": len(target_reader.pages) + len(source_reader.pages),
        **result,
    }
