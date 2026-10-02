from __future__ import annotations

import csv
import io
import json
from pathlib import Path
from typing import Any

from .commandport import _resolve

MAX_TEXT_CHARS = 128_000
MAX_ROWS = 500
MAX_COLS = 100
MAX_PAGES = 20


def _target(path: str) -> Path:
    target = _resolve(path, must_exist=True)
    if not target.is_file():
        raise IsADirectoryError(str(target))
    if target.is_symlink():
        raise PermissionError("symbolic-link document reads are not allowed")
    return target


def _cap(text: str, limit: int = MAX_TEXT_CHARS) -> tuple[str, bool]:
    if len(text) <= limit:
        return text, False
    return text[:limit] + "\n... Project Relay output truncated ...", True


def _read_csv(target: Path) -> dict[str, Any]:
    raw = target.read_text(encoding="utf-8-sig", errors="replace")
    rows: list[list[str]] = []
    delimiter = "\\t" if target.suffix.casefold() == ".tsv" else ","
    reader = csv.reader(io.StringIO(raw), delimiter=delimiter)
    for index, row in enumerate(reader):
        if index >= MAX_ROWS:
            break
        rows.append([str(value) for value in row[:MAX_COLS]])
    text = "\n".join("\t".join(row) for row in rows)
    text, truncated = _cap(text)
    return {"kind": "csv", "rows": rows, "text": text, "truncated": truncated or len(rows) >= MAX_ROWS}


def _read_json(target: Path) -> dict[str, Any]:
    value = json.loads(target.read_text(encoding="utf-8"))
    text = json.dumps(value, indent=2, ensure_ascii=False)
    text, truncated = _cap(text)
    return {"kind": "json", "text": text, "truncated": truncated}


def _read_pdf(target: Path, page_start: int, page_count: int) -> dict[str, Any]:
    from pypdf import PdfReader

    reader = PdfReader(str(target))
    start = max(0, int(page_start))
    count = max(1, min(int(page_count), MAX_PAGES))
    pages = []
    for index in range(start, min(len(reader.pages), start + count)):
        pages.append({"page": index + 1, "text": reader.pages[index].extract_text() or ""})
    text = "\n\n".join(f"[Page {p['page']}]\n{p['text']}" for p in pages)
    text, truncated = _cap(text)
    return {
        "kind": "pdf",
        "page_count": len(reader.pages),
        "pages_returned": len(pages),
        "text": text,
        "truncated": truncated or start + count < len(reader.pages),
    }


def _read_docx(target: Path) -> dict[str, Any]:
    from docx import Document

    doc = Document(str(target))
    lines: list[str] = [paragraph.text for paragraph in doc.paragraphs if paragraph.text]
    for table_index, table in enumerate(doc.tables[:50], start=1):
        lines.append(f"[Table {table_index}]")
        for row in table.rows[:MAX_ROWS]:
            lines.append("\t".join(cell.text for cell in row.cells[:MAX_COLS]))
    text, truncated = _cap("\n".join(lines))
    return {"kind": "docx", "text": text, "truncated": truncated}


def _read_xlsx(target: Path, sheet: str | None) -> dict[str, Any]:
    from openpyxl import load_workbook

    workbook = load_workbook(filename=str(target), read_only=True, data_only=False)
    try:
        if sheet:
            if sheet not in workbook.sheetnames:
                raise KeyError(f"worksheet not found: {sheet}")
            worksheets = [workbook[sheet]]
        else:
            worksheets = [workbook[name] for name in workbook.sheetnames[:20]]
        result_sheets = []
        text_parts = []
        truncated = False
        for worksheet in worksheets:
            rows = []
            for index, row in enumerate(worksheet.iter_rows(values_only=True)):
                if index >= MAX_ROWS:
                    truncated = True
                    break
                values = ["" if value is None else str(value) for value in row[:MAX_COLS]]
                rows.append(values)
            result_sheets.append({"name": worksheet.title, "rows": rows})
            text_parts.append(f"[Sheet: {worksheet.title}]")
            text_parts.extend("\t".join(row) for row in rows)
        text, cap_truncated = _cap("\n".join(text_parts))
        return {
            "kind": "xlsx",
            "sheet_names": list(workbook.sheetnames),
            "sheets": result_sheets,
            "text": text,
            "truncated": truncated or cap_truncated,
        }
    finally:
        workbook.close()


def read_document(
    path: str,
    sheet: str | None = None,
    page_start: int = 0,
    page_count: int = 10,
) -> dict[str, Any]:
    target = _target(path)
    suffix = target.suffix.casefold()
    if suffix in {".txt", ".md", ".py", ".js", ".ts", ".html", ".css", ".xml", ".log"}:
        text, truncated = _cap(target.read_text(encoding="utf-8", errors="replace"))
        result = {"kind": "text", "text": text, "truncated": truncated}
    elif suffix == ".json":
        result = _read_json(target)
    elif suffix in {".csv", ".tsv"}:
        result = _read_csv(target)
    elif suffix == ".pdf":
        result = _read_pdf(target, page_start, page_count)
    elif suffix == ".docx":
        result = _read_docx(target)
    elif suffix in {".xlsx", ".xlsm"}:
        result = _read_xlsx(target, sheet)
    else:
        raise ValueError("unsupported document type")
    return {"path": str(target), **result}


def search_document(
    path: str,
    query: str,
    case_sensitive: bool = False,
    sheet: str | None = None,
    max_results: int = 100,
) -> dict[str, Any]:
    needle = str(query or "")
    if not needle:
        raise ValueError("query is required")
    max_results = max(1, min(int(max_results), 500))
    doc = read_document(path, sheet=sheet, page_start=0, page_count=MAX_PAGES)
    haystack = str(doc.get("text") or "")
    target_needle = needle if case_sensitive else needle.casefold()
    matches = []
    for line_number, line in enumerate(haystack.splitlines(), start=1):
        candidate = line if case_sensitive else line.casefold()
        column = candidate.find(target_needle)
        if column < 0:
            continue
        matches.append({"line": line_number, "column": column + 1, "preview": line[:500]})
        if len(matches) >= max_results:
            break
    return {
        "path": doc["path"],
        "kind": doc["kind"],
        "query": needle,
        "case_sensitive": bool(case_sensitive),
        "count": len(matches),
        "truncated": len(matches) >= max_results or bool(doc.get("truncated")),
        "matches": matches,
    }
